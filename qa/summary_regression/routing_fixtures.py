"""Generate synthetic logo and scan layouts without application imports."""
import io
from pathlib import Path
import fitz
from PIL import Image
from docx import Document
from docx.shared import Inches
DATA = Path(__file__).resolve().parents[1] / "testdata" / "routing"

def fixture_documents():
    DATA.mkdir(exist_ok=True)
    stream = io.BytesIO()
    Image.new('RGB', (80, 30), 'navy').save(stream, format='PNG')
    logo = stream.getvalue()
    pdf = fitz.open()
    for _ in range(3):
        page = pdf.new_page()
        page.insert_image(fitz.Rect(36, 20, 116, 50), stream=logo)
        page.insert_text((36, 130), 'Ultrasound ordered; not performed.')
    (DATA / 'letterhead.pdf').write_bytes(pdf.tobytes())
    page = pdf.new_page()
    page.insert_image(fitz.Rect(0, 0, 500, 700), stream=logo)
    (DATA / 'mixed_scan.pdf').write_bytes(pdf.tobytes())
    pdf.close()
    for location in ('header', 'body', 'footer'):
        document = Document()
        paragraph = (document.sections[0].header.paragraphs[0] if location == 'header' else
                     document.sections[0].footer.paragraphs[0] if location == 'footer' else document.add_paragraph())
        paragraph.add_run().add_picture(io.BytesIO(logo), width=Inches(1))
        document.add_paragraph('Ultrasound ordered; not performed.')
        document.add_table(rows=1, cols=2).rows[0].cells[0].text = 'Diagnosis recorded'
        document.save(DATA / f'logo_{location}.docx')
    document = Document()
    document.add_paragraph('Report with attached clinical image')
    document.add_picture(io.BytesIO(logo), width=Inches(5))
    document.save(DATA / 'clinical_image.docx')

